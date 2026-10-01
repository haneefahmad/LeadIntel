"""
FastAPI Backend for Lead Intelligence System.
Exposes REST and SSE endpoints for search, analytics, live scraping, settings, and exports.
Mounted with modular frontend architecture.
"""

import asyncio
import base64
import csv
from datetime import datetime
import io
import json
import logging
import os
from pathlib import Path
import signal
import sys
import uuid
from typing import Any

from fastapi import FastAPI, BackgroundTasks, Query, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

# Ensure package roots are recognized
BACKEND_DIR = Path(__file__).resolve().parent.parent
ROOT_DIR = BACKEND_DIR.parent
APP_DIR = Path(__file__).resolve().parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

try:
    from backend.app import (
        config,
        database as db,
        scraper,
        website_checker,
        contact_enricher,
        exporter,
        geo_data,
        field_presets,
        apollo_client,
        enrichment_service,
    )
    scrape = scraper.scrape
    check_batch = website_checker.check_batch
    enrich_contacts_batch = contact_enricher.enrich_contacts_batch
    XLSXExporter = exporter.XLSXExporter
except (ImportError, ModuleNotFoundError) as e:
    if getattr(e, "name", None) not in ("backend", "backend.app", "app"):
        raise
    import config
    import database as db
    from scraper import scrape
    from website_checker import check_batch
    from contact_enricher import enrich_contacts_batch
    from exporter import XLSXExporter
    import geo_data
    import field_presets
    import apollo_client
    import enrichment_service

# Ensure DB is initialized
db.init_db()

logger = logging.getLogger("lead_intelligence_api")
logging.basicConfig(level=logging.INFO)

app = FastAPI(
    title="Lead Intelligence System API",
    description="Enterprise Lead Intelligence & Market Intelligence Platform",
    version="2.1.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# In-memory tracking for active and recent scrape jobs
active_jobs: dict[str, dict[str, Any]] = {}


# ── Pydantic Request Models ───────────────────────────────────────────────────

class ScrapeRequest(BaseModel):
    country: str = "United States"
    locations: list[str] = ["New York"]
    countries: list[str] = []
    states: list[str] = []
    cities: list[str] = []
    industries: list[str]
    max_records: int = 25
    added_by: str = "Web Dashboard"
    sheet_name: str | None = None
    is_new_sheet: bool = False


class SheetCreateRequest(BaseModel):
    name: str
    set_active: bool = True


class SheetSelectRequest(BaseModel):
    name: str = ""
    sheet: str = ""

    @property
    def target_name(self) -> str:
        return self.sheet or self.name


class SheetUploadRequest(BaseModel):
    sheet_name: str
    filename: str
    file_base64: str
    set_active: bool = True


class SingleCompanyExtractRequest(BaseModel):
    company_name: str
    city: str = ""
    country: str = "Saudi Arabia"
    domain: str = ""
    engines: list[str] = ["apify", "apollo"]
    sheet_name: str | None = None


class SettingsRequest(BaseModel):
    apify_token: str | None = None
    apollo_api_key: str | None = None


class DatabaseConfigRequest(BaseModel):
    database_url: str | None = None
    db_type: str | None = None
    host: str | None = None
    port: int | None = None
    username: str | None = None
    password: str | None = None
    database: str | None = None
    ssl_mode: str | None = None


class EnrichPreviewRequest(BaseModel):
    sheet_name: str | None = None
    engine: str  # "apify" | "apollo"
    selected_ids: list[str] | None = None


class EnrichStartRequest(BaseModel):
    sheet_name: str | None = None
    engine: str  # "apify" | "apollo"
    selected_ids: list[str] | None = None
    added_by: str = "Lead Intelligence Dashboard"
    run_id: str | None = None
    dataset_id: str | None = None


class RecordUpdateRequest(BaseModel):
    status: str | None = None
    pipeline_stage: str | None = None
    outreach_status: str | None = None
    assigned_to: str | None = None
    deal_value_sar: float | None = None
    notes: str | None = None
    tags: str | None = None
    next_action_date: str | None = None
    next_action_type: str | None = None
    email_engagement_status: str | None = None
    total_touches: int | None = None
    replies_received: int | None = None
    updated_by: str = "Web Dashboard"


class SuppressRequest(BaseModel):
    reason: str = "Customer Opt-out / PDPL Request"
    suppressed_by: str = "Web Dashboard"
    notes: str = ""


# ── Background Pipeline Runner ────────────────────────────────────────────────

async def _run_scrape_job(job_id: str, req: ScrapeRequest):
    target_sheet = config.clean_sheet_name(req.sheet_name or config.get_active_sheet())
    config.set_active_sheet(target_sheet)
    db.init_db(target_sheet)

    target_locations = req.cities if req.cities else req.locations
    if not target_locations:
        target_locations = ["New York"]

    fallback_country = req.country or (req.countries[0] if req.countries else "United States")

    job = active_jobs[job_id]
    job["sheet_name"] = target_sheet
    job["status"] = "running"
    job["started_at"] = datetime.now().isoformat()
    job["total_steps"] = len(target_locations) * len(req.industries)
    job["completed_steps"] = 0
    job["total_saved"] = 0
    job["total_cost"] = 0.0

    try:
        def log(msg: str, level: str = "info"):
            ts = datetime.now().strftime("%H:%M:%S")
            entry = {"time": ts, "message": msg, "level": level}
            job["logs"].append(entry)
            logger.info("[%s] %s", job_id, msg)

        log(f"Target Master Sheet: '{target_sheet}' ({config.DB_PATH})")
        log(f"Initiating mission: {len(target_locations)} location(s) × {len(req.industries)} industry(s)")

        total_runs = len(target_locations) * len(req.industries)
        run_num = 0

        for loc_entry in target_locations:
            if job.get("cancelled"):
                log("Mission cancelled by user.", "warning")
                return
            city, loc_country, display_loc = geo_data.resolve_location_details(loc_entry, fallback_country)
            for industry_key in req.industries:
                if job.get("cancelled"):
                    log("Mission cancelled by user.", "warning")
                    return
                run_num += 1
                industry_cfg = config.INDUSTRIES.get(industry_key, {})
                industry_name = industry_cfg.get("name", industry_key)

                job["current_location"] = display_loc
                job["current_industry"] = industry_name
                job["current_run"] = run_num
                job["stage"] = "Stage 1: Apify Google Maps Scrape"

                log(f"[{run_num}/{total_runs}] Starting scrape: {display_loc} — {industry_name}")

                queries = industry_cfg.get("queries") or [industry_name]
                all_raw: list[dict] = []
                t_start = datetime.now()

                for query in queries:
                    if job.get("cancelled"):
                        log("Mission cancelled by user.", "warning")
                        return
                    log(f"  Querying: '{query} in {city}' (max {req.max_records})")
                    batch: list[dict] = []
                    try:
                        async for record in scrape(
                            query=query,
                            city=city,
                            country=loc_country,
                            max_records=req.max_records,
                            industry_key=industry_key,
                            added_by=req.added_by,
                        ):
                            batch.append(record)
                    except Exception as err:
                        log(f"  [!] Scrape error: {err}", "error")
                        break

                    all_raw.extend(batch)
                    batch_cost = len(batch) * 0.011
                    job["total_cost"] += batch_cost
                    log(f"  Retrieved {len(batch)} places (${batch_cost:.2f})")

                if not all_raw:
                    log(f"  No places returned for {display_loc} / {industry_name}", "warning")
                    job["completed_steps"] += 1
                    continue

                # Stage 2: Deduplication & Freshness check
                job["stage"] = "Stage 2: Deduplication & Freshness Check"
                all_place_ids = [r.get("google_place_id", "") for r in all_raw if r.get("google_place_id")]
                fresh_cache = db.get_existing_enrichment(all_place_ids, max_age_days=30)

                dupes = 0
                records_needing_enrichment: list[dict] = []
                for rec in all_raw:
                    pid = rec.get("google_place_id", "")
                    if pid and db.record_exists(pid):
                        dupes += 1
                    if pid in fresh_cache:
                        stored = fresh_cache[pid]
                        for col in db.MASTER_COLUMNS:
                            if rec.get(col) in (None, "") and stored.get(col) not in (None, ""):
                                rec[col] = stored[col]
                    else:
                        records_needing_enrichment.append(rec)

                new_count = len(all_raw) - dupes
                cached_count = len(all_raw) - len(records_needing_enrichment)
                log(f"  Deduplication: {new_count} new, {dupes} existing ({cached_count} fresh skipped)")

                website_signals: dict[str, dict] = {}
                contact_data: dict[str, dict] = {}

                if records_needing_enrichment:
                    # Stage 3: Concurrent Website Checks (only for new/stale records)
                    job["stage"] = "Stage 3: Concurrent Website Audit"
                    has_websites = [r for r in records_needing_enrichment if r.get("Website_URL")]
                    log(f"  Auditing {len(has_websites)} websites via HTTPX ({cached_count} cached skipped)...")
                    html_cache: dict[str, str] = {}
                    website_signals = await check_batch(records_needing_enrichment, html_cache=html_cache)
                    log(f"  Website audit completed.")

                    # Stage 4: Contact Enrichment (only for new/stale records)
                    job["stage"] = "Stage 4: Website Contact Discovery"
                    log(f"  Scanning contacts ({len(records_needing_enrichment)} leads)...")
                    contact_data = await enrich_contacts_batch(records_needing_enrichment, html_cache=html_cache)
                    log(f"  Enrichment completed: {len(contact_data)} domains enriched.")

                # Stage 5: Save to master.db
                job["stage"] = "Stage 5: Persisting to Database"
                for record in all_raw:
                    pid = record.get("google_place_id", "")
                    ws = website_signals.get(pid, {})
                    for k, v in ws.items():
                        if not k.startswith("_"):
                            record[k] = v

                    contact = contact_data.get(pid, {})
                    for k, v in contact.items():
                        if not k.startswith("_"):
                            record[k] = v

                    record.pop("_raw_categories", None)

                saved, updated = db.upsert_records_batch(all_raw)
                job["total_saved"] += saved
                log(f"  Persisted {len(all_raw)} records to master.db ({saved} new, {updated} updated)")

                # Stage 6: Log run
                t_end = datetime.now()
                db.log_run({
                    "date": t_start.date().isoformat(),
                    "time_started": t_start.strftime("%H:%M"),
                    "time_completed": t_end.strftime("%H:%M"),
                    "query": ", ".join(queries),
                    "language": "en",
                    "city": city,
                    "industry": industry_name,
                    "max_records": req.max_records,
                    "records_returned": len(all_raw),
                    "duplicates": dupes,
                    "added_to_master": saved,
                    "cost_usd": round(len(all_raw) * 0.011, 3),
                    "quality_score": 9.0,
                    "run_by": req.added_by,
                    "run_type": "Web Mission",
                })

                job["completed_steps"] += 1

        # Auto export Excel after full job completes
        try:
            log(f"Regenerating {target_sheet}.xlsx file...")
            XLSXExporter().export(config.XLSX_PATH)
            log(f"{target_sheet}.xlsx successfully generated.")
        except Exception as e:
            log(f"Excel export notice: {e}", "warning")

        job["status"] = "completed"
        job["stage"] = "All Stages Complete"
        job["finished_at"] = datetime.now().isoformat()
        log(f"Mission finished! Total saved to DB: {job['total_saved']}")

    except Exception as e:
        job["status"] = "failed"
        job["error"] = str(e)
        job["logs"].append({
            "time": datetime.now().strftime("%H:%M:%S"),
            "message": f"[FATAL] Pipeline terminated unexpectedly: {e}",
            "level": "error"
        })
        logger.exception("Error in job %s", job_id)


class GeoStatesRequest(BaseModel):
    countries: list[str] = []


class GeoCitiesRequest(BaseModel):
    countries: list[str] = []
    states: list[str] = []


# ── REST API Endpoints ────────────────────────────────────────────────────────

@app.get("/api/geo/hierarchy")
def get_geo_hierarchy():
    """Returns the full hierarchical Country -> State -> Cities tree and country list."""
    return {
        "countries": geo_data.get_all_countries(),
        "hierarchy": geo_data.get_full_hierarchy(),
    }


@app.get("/api/geo/countries")
def get_geo_countries():
    """Returns a sorted list of all available countries."""
    return {"countries": geo_data.get_all_countries()}


@app.post("/api/geo/states")
def get_geo_states(req: GeoStatesRequest):
    """Returns states/provinces dynamically filtered for the given countries."""
    return {"states": geo_data.get_states_for_countries(req.countries)}


@app.post("/api/geo/cities")
def get_geo_cities(req: GeoCitiesRequest):
    """Returns cities dynamically filtered by countries and states."""
    return {"cities": geo_data.get_cities_for_selection(req.countries, req.states)}


@app.get("/api/fields")
def get_fields_catalog():
    """Returns the comprehensive 45-column schema catalog, categories, and presets."""
    return field_presets.get_field_catalog()


@app.get("/api/config")
def get_system_config(country: str = Query("", description="Optional country to filter presets")):
    """Returns industries, preset cities, depth options, and global hubs."""
    industries_list = [
        {
            "key": k,
            "name": v["name"],
            "query_count": len(v.get("queries", [])),
            "default_sub": v.get("default_sub", "General"),
            "queries": v.get("queries", []),
        }
        for k, v in config.INDUSTRIES.items()
    ]
    return {
        "industries": industries_list,
        "countries": geo_data.get_all_countries(),
        "preset_cities": config.get_preset_cities(country),
        "global_hubs": config.GLOBAL_BUSINESS_HUBS,
        "country_presets": config.COUNTRY_CITIES,
        "depth_presets": config.CRAWL_DEPTH_PRESETS,
        "database_path": str(config.DB_PATH),
        "xlsx_path": str(config.XLSX_PATH),
        "active_sheet": config.get_active_sheet(),
        "sheets": config.list_available_sheets(),
    }


@app.get("/api/sheets")
def get_sheets():
    """Lists all available master sheets and indicates the currently active sheet."""
    return {
        "active_sheet": config.get_active_sheet(),
        "sheets": config.list_available_sheets(),
    }


@app.post("/api/sheets")
def create_sheet(req: SheetCreateRequest):
    """Creates a new master sheet database and initial XLSX workbook."""
    clean = config.clean_sheet_name(req.name)
    db_path, xlsx_path = config.get_sheet_paths(clean)

    # Initialize the new database schema
    db.init_db(clean)

    # Export initial template XLSX
    try:
        XLSXExporter().export(str(xlsx_path))
    except Exception as exc:
        logger.warning("Could not create initial XLSX for %s: %s", clean, exc)

    if not req.set_active:
        config.set_active_sheet(config.DEFAULT_SHEET_NAME)

    return {
        "status": "success",
        "sheet": {
            "name": clean,
            "db_path": str(db_path),
            "xlsx_path": str(xlsx_path),
            "is_active": config.get_active_sheet() == clean,
        },
        "active_sheet": config.get_active_sheet(),
        "sheets": config.list_available_sheets(),
        "message": f"Sheet '{clean}' created successfully.",
    }


@app.post("/api/sheets/select")
def select_sheet(req: SheetSelectRequest):
    """Switches the active master sheet for the dashboard."""
    target = req.target_name
    clean = config.clean_sheet_name(target)
    db_path, _ = config.get_sheet_paths(clean)
    if not db_path.exists() and clean != config.DEFAULT_SHEET_NAME:
        raise HTTPException(status_code=404, detail=f"Sheet '{clean}' not found.")

    config.set_active_sheet(clean)
    db.init_db(clean)

    return {
        "status": "success",
        "active_sheet": clean,
        "sheets": config.list_available_sheets(),
        "message": f"Active sheet switched to '{clean}'.",
    }


@app.post("/api/sheets/upload")
def upload_datasheet(req: SheetUploadRequest):
    """
    Accepts an uploaded Excel (.xlsx, .xls) or CSV file in base64 format,
    maps columns into canonical master schema, creates or updates a dedicated sheet database,
    and returns sheet metadata ready for immediate viewing and on-demand extraction/enrichment.
    """
    if not req.file_base64:
        raise HTTPException(status_code=400, detail="No file data provided.")

    filename = req.filename or "uploaded_sheet.xlsx"
    ext = Path(filename).suffix.lower()
    if ext not in (".xlsx", ".csv", ".xls"):
        raise HTTPException(status_code=400, detail="Only Excel (.xlsx, .xls) and CSV (.csv) files are supported.")

    try:
        b64_str = req.file_base64
        if "," in b64_str:
            b64_str = b64_str.split(",", 1)[1]
        file_bytes = base64.b64decode(b64_str)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid base64 payload: {e}")

    raw_records: list[dict[str, Any]] = []

    if ext in (".xlsx", ".xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                raise HTTPException(status_code=400, detail="Uploaded Excel file is empty or missing headers.")
            headers = [str(h).strip() if h is not None else f"Column_{i+1}" for i, h in enumerate(rows[0])]
            for row in rows[1:]:
                if not any(row):
                    continue
                rec = {}
                for h, val in zip(headers, row):
                    if val is not None:
                        rec[h] = str(val).strip() if isinstance(val, (int, float, str)) else val
                if rec:
                    raw_records.append(rec)
        except HTTPException:
            raise
        except Exception as e:
            logger.exception("Excel parsing failed: %s", e)
            raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {e}")

    elif ext == ".csv":
        text_content = ""
        for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                text_content = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if not text_content:
            raise HTTPException(status_code=400, detail="Could not decode CSV file. Please ensure it is saved in UTF-8.")

        try:
            reader = csv.DictReader(io.StringIO(text_content))
            for row in reader:
                clean_row = {str(k).strip(): str(v).strip() for k, v in row.items() if k is not None and v is not None}
                if any(clean_row.values()):
                    raw_records.append(clean_row)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse CSV file: {e}")

    if not raw_records:
        raise HTTPException(status_code=400, detail="No readable data rows found in the uploaded file.")

    default_name = Path(filename).stem
    sheet_raw = req.sheet_name.strip() if req.sheet_name else default_name
    clean_sheet = config.clean_sheet_name(sheet_raw)
    db_path, xlsx_path = config.get_sheet_paths(clean_sheet)

    db.init_db(clean_sheet)

    normalized_list = []
    now_str = datetime.now().strftime("%Y-%m-%d")

    for i, raw in enumerate(raw_records):
        norm: dict[str, Any] = {}
        unmapped: list[str] = []

        for key, val in raw.items():
            if val is None or str(val).strip() == "":
                continue
            k_clean = str(key).strip()
            v_clean = str(val).strip()

            canonical = db.COLUMN_ALIASES.get(k_clean)
            if not canonical:
                for alias_k, canon_v in db.COLUMN_ALIASES.items():
                    if alias_k.lower() == k_clean.lower():
                        canonical = canon_v
                        break

            if canonical and canonical in db.MASTER_COLUMNS:
                norm[canonical] = v_clean
            else:
                unmapped.append(f"{k_clean}: {v_clean}")

        if not norm.get("Company_Name") and norm.get("Website_URL"):
            domain = apollo_client.extract_clean_domain(norm["Website_URL"])
            if domain:
                norm["Company_Name"] = domain.split(".")[0].capitalize()

        if not norm.get("Company_Name"):
            if not norm.get("Website_URL") and not norm.get("Primary_Phone") and not norm.get("DM_Full_Name"):
                continue
            norm["Company_Name"] = f"Lead-{i+1}"

        if unmapped and not norm.get("CRM_Notes"):
            norm["CRM_Notes"] = " | ".join(unmapped[:6])

        norm["Added_By"] = "Datasheet Upload"
        norm["Date_Added"] = norm.get("Date_Added") or now_str
        norm["Lead_Status"] = norm.get("Lead_Status") or "New"
        norm["Pipeline_Stage"] = norm.get("Pipeline_Stage") or "Lead"
        norm["Outreach_Status"] = norm.get("Outreach_Status") or "Uncontacted"
        if norm.get("Website_URL"):
            norm["Has_Website"] = "Yes"

        normalized_list.append(norm)

    if not normalized_list:
        raise HTTPException(status_code=400, detail="No valid business records could be extracted from the uploaded file.")

    inserted, updated = db.upsert_records_batch(normalized_list, commit=True, sheet_name=clean_sheet)

    try:
        prev_sheet = config.get_active_sheet()
        config.set_active_sheet(clean_sheet)
        XLSXExporter().export(str(xlsx_path))
        if not req.set_active:
            config.set_active_sheet(prev_sheet)
    except Exception as exc:
        logger.warning("Could not export initial XLSX for %s: %s", clean_sheet, exc)

    if req.set_active:
        config.set_active_sheet(clean_sheet)

    return {
        "status": "success",
        "sheet": {
            "name": clean_sheet,
            "db_path": str(db_path),
            "xlsx_path": str(xlsx_path),
            "is_active": config.get_active_sheet() == clean_sheet,
        },
        "rows_processed": len(raw_records),
        "inserted": inserted,
        "updated": updated,
        "active_sheet": config.get_active_sheet(),
        "sheets": config.list_available_sheets(),
        "message": f"Successfully uploaded and mapped {len(normalized_list)} records into sheet '{clean_sheet}'.",
    }


@app.delete("/api/sheets/{sheet_name}")
def delete_sheet_endpoint(sheet_name: str):
    """Permanently deletes a custom master sheet and its associated database/export files."""
    clean = config.clean_sheet_name(sheet_name)
    if clean == config.DEFAULT_SHEET_NAME:
        raise HTTPException(
            status_code=400,
            detail="The default 'MasterDB' sheet is protected and cannot be deleted.",
        )

    # Check if a scrape job is actively running for this sheet
    for job_id, job in active_jobs.items():
        if job.get("status") == "running":
            job_sheet = config.clean_sheet_name(job.get("sheet_name") or config.DEFAULT_SHEET_NAME)
            if job_sheet == clean:
                raise HTTPException(
                    status_code=409,
                    detail=f"Cannot delete sheet '{clean}' while scrape mission #{job_id} is actively running on it.",
                )

    # Close any active db connection to release file locks
    db.close_connection()

    try:
        res = config.delete_sheet(clean)
        # Ensure default sheet is initialized
        db.init_db(config.DEFAULT_SHEET_NAME)
        return {
            "status": "success",
            "message": f"Master sheet '{clean}' was successfully deleted.",
            "deleted_sheet": clean,
            "active_sheet": config.get_active_sheet(),
            "sheets": config.list_available_sheets(),
            "details": res,
        }
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as exc:
        logger.exception("Error deleting sheet %s: %s", clean, exc)
        raise HTTPException(status_code=500, detail=f"Failed to delete sheet: {exc}")


@app.get("/api/stats")
def get_database_stats(sheet: str = Query("", description="Optional master sheet name")):
    """Returns real-time KPIs and distribution breakdown."""
    if sheet:
        clean = config.clean_sheet_name(sheet)
        config.set_active_sheet(clean)
        db.init_db(clean)
    return db.get_stats()


@app.get("/api/records")
def get_records(
    sheet: str = Query("", description="Optional master sheet name"),
    q: str = Query("", description="Keyword search across name, address, email, phone"),
    city: str = Query("", description="Filter by city"),
    industry: str = Query("", description="Filter by industry"),
    status: str = Query("", description="Filter by status"),
    has_website: str = Query("", description="Yes or No"),
    has_email: str = Query("", description="Yes or No"),
    page: int = Query(1, ge=1),
    limit: int = Query(25, ge=5, le=200),
):
    """Filterable, paginated lead records backed by indexed SQL queries."""
    if sheet:
        clean = config.clean_sheet_name(sheet)
        config.set_active_sheet(clean)
        db.init_db(clean)
    page_records, total_count = db.query_records(
        q=q,
        city=city,
        industry=industry,
        status=status,
        has_website=has_website,
        has_email=has_email,
        page=page,
        limit=limit,
    )
    total_pages = max(1, (total_count + limit - 1) // limit)
    return {
        "records": page_records,
        "total": total_count,
        "page": page,
        "pages": total_pages,
        "limit": limit,
    }


@app.get("/api/records/{record_id}")
def get_record_detail(record_id: str):
    """Retrieve full details of a specific company."""
    record = db.get_record_by_id(record_id)
    if not record:
        raise HTTPException(status_code=404, detail="Record not found")
    return record


@app.patch("/api/records/{record_id}")
def update_record(record_id: str, req: RecordUpdateRequest):
    """Update CRM fields for a lead record."""
    field_map = {
        "status": "Status",
        "pipeline_stage": "Pipeline_Stage",
        "outreach_status": "Outreach_Status",
        "assigned_to": "Assigned_To",
        "deal_value_sar": "Deal_Value_SAR",
        "notes": "Notes",
        "tags": "Tags",
        "next_action_date": "Next_Action_Date",
        "next_action_type": "Next_Action_Type",
        "email_engagement_status": "Email_Engagement_Status",
        "total_touches": "Total_Touches",
        "replies_received": "Replies_Received",
    }
    updates = {}
    for req_field, db_col in field_map.items():
        val = getattr(req, req_field, None)
        if val is not None:
            updates[db_col] = val

    updated_rec = db.update_record_crm(record_id, updates, updated_by=req.updated_by)
    if not updated_rec:
        raise HTTPException(status_code=404, detail="Record not found")
    return {"status": "success", "record": updated_rec}


@app.post("/api/records/{record_id}/suppress")
def suppress_lead(record_id: str, req: SuppressRequest = SuppressRequest()):
    """Suppress a lead for PDPL/GDPR compliance and mark as Disqualified/Do Not Contact."""
    success = db.suppress_record(
        record_id=record_id,
        reason=req.reason,
        suppressed_by=req.suppressed_by,
        notes=req.notes,
    )
    if not success:
        raise HTTPException(status_code=404, detail="Record not found")
    return {"status": "success", "message": f"Lead {record_id} successfully suppressed."}


@app.post("/api/companies/extract")
async def extract_single_company(req: SingleCompanyExtractRequest):
    """
    Extracts high-fidelity company details using Apify (Google Maps/Places),
    Apollo (B2B contacts/executive leadership), or both, based on user selection.
    Saves or non-destructively merges the result into the target master sheet.
    """
    company_name = req.company_name.strip()
    if not company_name:
        raise HTTPException(status_code=400, detail="Company name is required.")

    target_sheet = config.clean_sheet_name(req.sheet_name or config.get_active_sheet())
    config.set_active_sheet(target_sheet)
    db.init_db(target_sheet)

    selected_engines = [e.lower().strip() for e in (req.engines or ["apify", "apollo"])]
    if not selected_engines:
        raise HTTPException(status_code=400, detail="At least one extraction engine (Apify or Apollo) must be selected.")

    # Check for existing record in database to update non-destructively
    existing_records, _ = db.query_records(q=company_name, limit=5)
    matched_record = None
    for r in existing_records:
        if (r.get("Company_Name") or "").strip().lower() == company_name.lower():
            matched_record = dict(r)
            break

    record: dict[str, Any] = dict(matched_record) if matched_record else {}
    record["Company_Name"] = record.get("Company_Name") or company_name
    if req.city and not record.get("City"):
        record["City"] = req.city.strip()
    if req.country and not record.get("Country"):
        record["Country"] = req.country.strip()
    if req.domain and not record.get("Website_URL"):
        clean_d = apollo_client.extract_clean_domain(req.domain)
        record["Website_URL"] = req.domain.strip() if req.domain.startswith("http") else f"https://{clean_d or req.domain.strip()}"
        record["Has_Website"] = "Yes"

    engine_notes: list[str] = []

    # ── Engine 1: Apify Google Maps / Places ──────────────────────────────────
    if "apify" in selected_engines:
        try:
            apify_data = await scraper.scrape_single_company(
                company_name=company_name,
                city=record.get("City") or req.city or "",
                country=record.get("Country") or req.country or "Saudi Arabia",
            )
            if apify_data:
                for k, v in apify_data.items():
                    if v and enrichment_service.is_field_empty(record.get(k), k):
                        record[k] = v
                engine_notes.append("Apify (Google Maps) firmographics retrieved")
            else:
                engine_notes.append("Apify (Google Maps) found no matching place")
        except PermissionError as pe:
            if selected_engines == ["apify"]:
                raise HTTPException(status_code=400, detail=str(pe))
            logger.warning("Apify single scrape permission error: %s", pe)
            engine_notes.append("Apify skipped (API token missing)")
        except Exception as e:
            logger.warning("Apify single scrape error: %s", e)
            engine_notes.append(f"Apify error: {e}")

    # ── Engine 2: Apollo.io Decision Maker & Contact Enrichment ───────────────
    if "apollo" in selected_engines:
        ap_client = apollo_client.ApolloClient()
        is_ready, ap_msg = ap_client.is_configured()
        if is_ready:
            try:
                apollo_fields = await ap_client.enrich_lead(record, enrich_firmographics=True)
                if apollo_fields:
                    for k, v in apollo_fields.items():
                        if v and enrichment_service.is_field_empty(record.get(k), k):
                            record[k] = v
                    engine_notes.append("Apollo.io executive contact & email enriched")
                else:
                    engine_notes.append("Apollo.io found no additional verified contact match")
            except Exception as e:
                logger.warning("Apollo enrichment error: %s", e)
                engine_notes.append(f"Apollo error: {e}")
        else:
            if selected_engines == ["apollo"]:
                raise HTTPException(status_code=400, detail=ap_msg)
            engine_notes.append(f"Apollo skipped ({ap_msg})")

    # Hygiene & scoring
    if not record.get("Record_ID"):
        record["Record_ID"] = f"REC-{uuid.uuid4().hex[:8].upper()}"
    if not record.get("Date_Added"):
        record["Date_Added"] = datetime.now().strftime("%Y-%m-%d")
    record["Added_By"] = record.get("Added_By") or "Direct Extraction"
    record["Lead_Status"] = record.get("Lead_Status") or "New"
    record["Pipeline_Stage"] = record.get("Pipeline_Stage") or "Lead"
    record["Outreach_Status"] = record.get("Outreach_Status") or "Uncontacted"
    if record.get("Website_URL") and not record.get("Has_Website"):
        record["Has_Website"] = "Yes"

    # Save non-destructively to database
    is_new, final_id = db.upsert_record(record, sheet_name=target_sheet)
    record["Record_ID"] = final_id

    # Log run
    try:
        now_time = datetime.now().strftime("%H:%M")
        db.log_run(
            {
                "run_type": "Single Company Extraction",
                "query": company_name,
                "city": record.get("City") or req.city or "Global",
                "industry": record.get("Primary_Industry") or "General",
                "records_returned": 1,
                "added_to_master": 1 if is_new else 0,
                "cost_usd": 0.01 if "apify" in selected_engines else 0.0,
                "time_started": now_time,
                "time_completed": now_time,
                "notes": f"Engines: {', '.join(selected_engines)}",
            },
            sheet_name=target_sheet,
        )
    except Exception as e:
        logger.warning("Failed to log single extraction run: %s", e)

    return {
        "status": "success",
        "action": "created" if is_new else "updated",
        "message": f"Successfully processed '{record.get('Company_Name')}' ({', '.join(engine_notes)}).",
        "record": record,
        "sheet_name": target_sheet,
        "engine_notes": engine_notes,
    }


@app.post("/api/scrape/start")
async def start_scrape(req: ScrapeRequest, background_tasks: BackgroundTasks):
    """Launch an autonomous multi-stage lead intelligence mission."""
    token = (config.APIFY_API_TOKEN or "").strip()
    if not token or token.lower() in ("your_apify_token_here", "your_token_here"):
        raise HTTPException(
            status_code=400,
            detail="Apify API Token is not configured. Please open Settings in the dashboard or configure APIFY_API_TOKEN in .env."
        )

    if not req.locations:
        raise HTTPException(status_code=400, detail="At least one location must be provided.")
    if not req.industries:
        raise HTTPException(status_code=400, detail="At least one industry must be selected.")

    job_id = str(uuid.uuid4())[:8]
    active_jobs[job_id] = {
        "id": job_id,
        "status": "queued",
        "params": req.dict(),
        "stage": "Initializing",
        "current_location": "",
        "current_industry": "",
        "completed_steps": 0,
        "total_steps": len(req.locations) * len(req.industries),
        "total_saved": 0,
        "total_cost": 0.0,
        "logs": [],
        "created_at": datetime.now().isoformat(),
    }

    background_tasks.add_task(_run_scrape_job, job_id, req)
    return {"job_id": job_id, "status": "queued"}


@app.get("/api/scrape/status/{job_id}")
def get_scrape_status(job_id: str):
    """Poll scrape job progress and recent logs."""
    if job_id not in active_jobs:
        raise HTTPException(status_code=404, detail="Job ID not found")
    return active_jobs[job_id]


@app.get("/api/scrape/stream/{job_id}")
async def stream_scrape_logs(job_id: str):
    """Server-Sent Events stream for real-time console logs."""
    if job_id not in active_jobs:
        raise HTTPException(status_code=404, detail="Job ID not found")

    async def event_generator():
        last_log_idx = 0
        while True:
            job = active_jobs.get(job_id)
            if not job:
                break

            logs = job.get("logs", [])
            if len(logs) > last_log_idx:
                new_logs = logs[last_log_idx:]
                last_log_idx = len(logs)
                payload = {
                    "status": job.get("status"),
                    "stage": job.get("stage"),
                    "completed_steps": job.get("completed_steps"),
                    "total_steps": job.get("total_steps"),
                    "total_saved": job.get("total_saved"),
                    "total_cost": job.get("total_cost"),
                    "current_location": job.get("current_location"),
                    "current_industry": job.get("current_industry"),
                    "new_logs": new_logs,
                }
                yield f"data: {json.dumps(payload)}\n\n"

            if job.get("status") in ("completed", "failed"):
                # One final message
                payload = {
                    "status": job.get("status"),
                    "stage": job.get("stage"),
                    "completed_steps": job.get("completed_steps"),
                    "total_steps": job.get("total_steps"),
                    "total_saved": job.get("total_saved"),
                    "total_cost": job.get("total_cost"),
                    "new_logs": [],
                    "finished": True,
                }
                yield f"data: {json.dumps(payload)}\n\n"
                break

            await asyncio.sleep(0.8)

    return StreamingResponse(event_generator(), media_type="text/event-stream")


@app.get("/api/logs")
def get_run_history():
    """Returns execution run logs."""
    return db.get_run_log()


@app.get("/api/export")
def download_excel_export(
    sheet: str = Query("", description="Optional sheet name"),
    columns: str = Query("", description="Comma-separated column IDs"),
):
    """Generates and serves the latest MasterDB.xlsx file for the active or specified sheet, with optional column filtering."""
    target_sheet = config.clean_sheet_name(sheet or config.get_active_sheet())
    _, xlsx_path = config.get_sheet_paths(target_sheet)

    prev_sheet = config.get_active_sheet()
    if target_sheet != prev_sheet:
        config.set_active_sheet(target_sheet)
        db.init_db(target_sheet)

    col_list = [c.strip() for c in columns.split(",") if c.strip()] if columns else None

    if col_list:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        custom_path = config.DATA_DIR / f"{target_sheet}_custom_{timestamp}.xlsx"
        try:
            XLSXExporter().export(str(custom_path), columns=col_list)
            export_file = custom_path
            download_name = f"{target_sheet}_custom_{timestamp}.xlsx"
        except Exception as e:
            logger.error("Error generating custom Excel export for %s: %s", target_sheet, e)
            export_file = xlsx_path
            download_name = f"{target_sheet}.xlsx"
    else:
        try:
            XLSXExporter().export(str(xlsx_path))
            export_file = xlsx_path
            download_name = f"{target_sheet}.xlsx" if target_sheet != config.DEFAULT_SHEET_NAME else "Lead_Intelligence_MasterDB.xlsx"
        except Exception as e:
            logger.error("Error generating Excel export for %s: %s", target_sheet, e)
            export_file = xlsx_path
            download_name = f"{target_sheet}.xlsx"

    if target_sheet != prev_sheet:
        config.set_active_sheet(prev_sheet)

    if not Path(export_file).exists():
        raise HTTPException(status_code=404, detail=f"Export file for sheet '{target_sheet}' could not be generated.")

    return FileResponse(
        path=str(export_file),
        filename=download_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/api/export/csv")
def download_csv_export(
    sheet: str = Query("", description="Optional sheet name"),
    columns: str = Query("", description="Comma-separated column IDs"),
    q: str = Query("", description="Keyword search across name, address, email, phone"),
    city: str = Query("", description="Filter by city"),
    industry: str = Query("", description="Filter by industry"),
    status: str = Query("", description="Filter by status"),
    has_website: str = Query("", description="Yes or No"),
    has_email: str = Query("", description="Yes or No"),
):
    """Generates and streams a standard CRM-compatible CSV export of lead records."""
    target_sheet = config.clean_sheet_name(sheet or config.get_active_sheet())
    prev_sheet = config.get_active_sheet()
    if target_sheet != prev_sheet:
        config.set_active_sheet(target_sheet)
        db.init_db(target_sheet)

    records, _ = db.query_records(
        q=q,
        city=city,
        industry=industry,
        status=status,
        has_website=has_website,
        has_email=has_email,
        page=1,
        limit=None,
    )

    if target_sheet != prev_sheet:
        config.set_active_sheet(prev_sheet)

    col_ids = [c.strip() for c in columns.split(",") if c.strip()] if columns else None
    if col_ids:
        valid_cols = [c for c in col_ids if c in db.MASTER_COLUMNS]
    else:
        valid_cols = list(db.MASTER_COLUMNS)

    csv_columns = [(col, field_presets.get_column_label(col)) for col in valid_cols]

    def iter_csv():
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([label for _, label in csv_columns])
        yield b"\xef\xbb\xbf" + output.getvalue().encode("utf-8")
        output.seek(0)
        output.truncate(0)

        for rec in records:
            row = [str(rec.get(key) if rec.get(key) is not None else "") for key, _ in csv_columns]
            writer.writerow(row)
            yield output.getvalue().encode("utf-8")
            output.seek(0)
            output.truncate(0)

    filename = f"{target_sheet}_leads_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return StreamingResponse(
        iter_csv(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/settings")
def get_settings():
    """Masked API key status and health verification."""
    apify = (config.APIFY_API_TOKEN or "").strip()
    apollo = (getattr(config, "APOLLO_API_KEY", "") or "").strip()

    def mask(key: str) -> str:
        if not key or key.lower() in ("your_apify_token_here", "your_key_here", "your_apollo_key_here"):
            return ""
        if len(key) <= 8:
            return "••••••••"
        return key[:6] + "••••••••" + key[-4:]

    # Apollo configuration validator
    ap_client = apollo_client.ApolloClient(api_key=apollo)
    is_apollo_configured, apollo_msg = ap_client.is_configured()

    return {
        "apify_configured": bool(apify and apify.lower() not in ("your_apify_token_here", "your_token_here")),
        "apify_masked": mask(apify),
        "apollo_configured": is_apollo_configured,
        "apollo_masked": mask(apollo) if is_apollo_configured else "",
        "apollo_status": apollo_msg,
    }


@app.post("/api/settings")
def update_settings(req: SettingsRequest):
    """Save new API keys directly to .env files (root and backend)."""
    targets = [ROOT_DIR / ".env", BACKEND_DIR / ".env"]

    for env_path in targets:
        existing_lines = []
        if env_path.exists():
            existing_lines = env_path.read_text(encoding="utf-8").splitlines()

        new_dict: dict[str, str] = {}
        for line in existing_lines:
            line_s = line.strip()
            if "=" in line_s and not line_s.startswith("#"):
                k, v = line_s.split("=", 1)
                k_clean = k.strip()
                new_dict[k_clean] = v.strip()

        if req.apify_token is not None and req.apify_token.strip():
            new_dict["APIFY_API_TOKEN"] = req.apify_token.strip()
            config.APIFY_API_TOKEN = req.apify_token.strip()

        if req.apollo_api_key is not None and req.apollo_api_key.strip():
            new_dict["APOLLO_API_KEY"] = req.apollo_api_key.strip()
            config.APOLLO_API_KEY = req.apollo_api_key.strip()

        # Remove LINKEDIN_COOKIE if present in new_dict
        new_dict.pop("LINKEDIN_COOKIE", None)

        out = "\n".join(f"{k}={v}" for k, v in new_dict.items()) + "\n"
        env_path.write_text(out, encoding="utf-8")

    return {"status": "success", "message": "API keys updated successfully."}


@app.post("/api/settings/test-apollo")
async def test_apollo_endpoint(req: dict[str, Any] | None = None):
    """Test live connectivity and authentication with Apollo.io API."""
    key = None
    if req and isinstance(req, dict):
        key = req.get("apollo_api_key")
    client = apollo_client.ApolloClient(api_key=key)
    success, message, details = await client.test_connection()
    return {"success": success, "message": message, "details": details}


# ── Universal Database Connection Endpoints ───────────────────────────────────

def _resolve_db_url(req: DatabaseConfigRequest) -> str:
    raw = (req.database_url or "").strip()
    if raw:
        return raw
    db_type = (req.db_type or "sqlite").lower().strip()
    if db_type == "sqlite":
        return f"sqlite:///{config.DB_PATH}"

    user = (req.username or "").strip()
    pwd = (req.password or "").strip()
    host = (req.host or "localhost").strip()
    port = req.port
    dbname = (req.database or "leadintel").strip()
    ssl = (req.ssl_mode or "").strip()

    auth = f"{user}:{pwd}@" if user or pwd else ""
    ssl_param = f"?sslmode={ssl}" if ssl else ""

    if db_type == "postgresql":
        port_str = f":{port}" if port else ":5432"
        return f"postgresql+psycopg://{auth}{host}{port_str}/{dbname}{ssl_param}"
    elif db_type in ("mysql", "mariadb"):
        port_str = f":{port}" if port else ":3306"
        return f"mysql+pymysql://{auth}{host}{port_str}/{dbname}{ssl_param}"
    elif db_type == "mssql":
        port_str = f":{port}" if port else ":1433"
        return f"mssql+pyodbc://{auth}{host}{port_str}/{dbname}{ssl_param}"

    return raw


@app.get("/api/database/config")
def get_database_config():
    """Returns the current database connection configuration and health status."""
    url = config.get_database_url()
    db_type = config.get_database_type()
    masked = config.mask_database_url()
    is_sqlite = config.is_sqlite()
    active_sheet = config.get_active_sheet()

    ok, msg, details = db.test_connection(url)

    return {
        "database_type": db_type,
        "database_url_masked": masked,
        "is_sqlite": is_sqlite,
        "active_sheet": active_sheet,
        "connected": ok,
        "connection_message": msg,
        "details": details,
    }


@app.post("/api/database/test")
def test_database_endpoint(req: DatabaseConfigRequest):
    """Tests live connection to any specified database URL or parameters without saving."""
    target_url = _resolve_db_url(req)
    if not target_url:
        raise HTTPException(status_code=400, detail="Database URL or parameters are required.")
    ok, msg, details = db.test_connection(target_url)
    return {
        "success": ok,
        "message": msg,
        "details": details,
        "target_masked": config.mask_database_url(target_url),
    }


@app.post("/api/database/config")
def update_database_config(req: DatabaseConfigRequest):
    """Saves new database connection to .env, tests connection, and initializes schemas."""
    target_url = _resolve_db_url(req)
    if not target_url:
        raise HTTPException(status_code=400, detail="Database URL or parameters are required.")

    # 1. Test connection before committing
    ok, msg, details = db.test_connection(target_url)
    if not ok:
        raise HTTPException(status_code=400, detail=f"Could not connect to database: {msg}")

    # 2. Update memory state
    config.set_database_url(target_url)

    # 3. Save DATABASE_URL directly to .env files
    targets = [ROOT_DIR / ".env", BACKEND_DIR / ".env"]
    for env_path in targets:
        existing_lines = []
        if env_path.exists():
            existing_lines = env_path.read_text(encoding="utf-8").splitlines()

        new_dict: dict[str, str] = {}
        for line in existing_lines:
            line_s = line.strip()
            if "=" in line_s and not line_s.startswith("#"):
                k, v = line_s.split("=", 1)
                new_dict[k.strip()] = v.strip()

        new_dict["DATABASE_URL"] = target_url
        out = "\n".join(f"{k}={v}" for k, v in new_dict.items()) + "\n"
        env_path.write_text(out, encoding="utf-8")

    # 4. Refresh connection pool and initialize tables
    db.close_connection()
    db.init_db()

    return {
        "status": "success",
        "message": f"Connected and initialized {config.get_database_type().upper()} database successfully.",
        "database_type": config.get_database_type(),
        "database_url_masked": config.mask_database_url(),
        "is_sqlite": config.is_sqlite(),
    }


@app.post("/api/database/reset-sqlite")
def reset_database_to_sqlite():
    """Resets the active database connection to local SQLite master.db."""
    config.set_database_url("")

    targets = [ROOT_DIR / ".env", BACKEND_DIR / ".env"]
    for env_path in targets:
        if env_path.exists():
            existing_lines = env_path.read_text(encoding="utf-8").splitlines()
            new_lines = [l for l in existing_lines if not l.strip().startswith("DATABASE_URL=")]
            env_path.write_text("\n".join(new_lines) + "\n", encoding="utf-8")

    db.close_connection()
    db.init_db()

    return {
        "status": "success",
        "message": "Reset to local SQLite database successfully.",
        "database_type": "sqlite",
        "database_url_masked": config.mask_database_url(),
        "is_sqlite": True,
    }


# ── Modular Lead Enrichment Endpoints ─────────────────────────────────────────

@app.post("/api/enrich/preview")
def preview_enrichment_endpoint(req: EnrichPreviewRequest):
    """Calculates missing fields, eligible records, and cost/credits preview."""
    try:
        preview = enrichment_service.preview_enrichment(
            sheet_name=req.sheet_name or config.get_active_sheet(),
            engine=req.engine,
            selected_ids=req.selected_ids,
        )
        return preview
    except Exception as exc:
        logger.exception("Error calculating enrichment preview: %s", exc)
        raise HTTPException(status_code=500, detail=str(exc))


@app.post("/api/enrich/start")
async def start_enrichment_endpoint(req: EnrichStartRequest):
    """Initiates an asynchronous background enrichment job streamed via SSE."""
    target_sheet = config.clean_sheet_name(req.sheet_name or config.get_active_sheet())
    job_id = f"enr_{uuid.uuid4().hex[:8]}"

    active_jobs[job_id] = {
        "job_id": job_id,
        "sheet_name": target_sheet,
        "status": "queued",
        "engine": req.engine,
        "stage": f"Queued {req.engine.title()} Enrichment",
        "started_at": datetime.now().isoformat(),
        "finished_at": None,
        "completed_steps": 0,
        "total_steps": 0,
        "total_saved": 0,
        "total_cost": 0.0,
        "current_location": "",
        "current_industry": "",
        "logs": [],
        "error": None,
    }

    asyncio.create_task(
        enrichment_service.run_enrichment_job(
            job_id=job_id,
            sheet_name=target_sheet,
            engine=req.engine,
            selected_ids=req.selected_ids,
            active_jobs=active_jobs,
            run_id=req.run_id,
            dataset_id=req.dataset_id,
        )
    )

    return {
        "job_id": job_id,
        "status": "started",
        "sheet_name": target_sheet,
        "engine": req.engine,
        "message": f"{req.engine.title()} enrichment job started successfully.",
    }


# ── Mission and Server Control Endpoints ─────────────────────────────────────

@app.post("/api/scrape/jobs/{job_id}/stop")
def stop_scrape_job(job_id: str):
    """Signals an active scraping or enrichment mission to abort immediately."""
    if job_id not in active_jobs:
        raise HTTPException(status_code=404, detail="Job ID not found")
    job = active_jobs[job_id]
    job["cancelled"] = True
    job["status"] = "failed"
    job["stage"] = "Mission stopped by user"
    job["finished_at"] = datetime.now().isoformat()
    ts = datetime.now().strftime("%H:%M:%S")
    job.setdefault("logs", []).append({
        "time": ts,
        "message": "⏹ Mission was manually stopped by user.",
        "level": "warning"
    })
    return {"status": "stopped", "job_id": job_id, "message": "Mission cancellation signal sent."}


@app.post("/api/server/shutdown")
async def shutdown_server():
    """Gracefully shuts down the Uvicorn/FastAPI server and releases port 8000."""
    logger.info("Server shutdown requested via API.")

    async def _delayed_shutdown():
        await asyncio.sleep(0.6)
        pid = os.getpid()
        ppid = os.getppid()
        if ppid > 1:
            try:
                os.kill(ppid, signal.SIGTERM)
            except Exception:
                pass
        try:
            os.kill(pid, signal.SIGTERM)
        except Exception:
            pass
        await asyncio.sleep(0.3)
        try:
            os._exit(0)
        except Exception:
            pass

    asyncio.create_task(_delayed_shutdown())
    return {
        "status": "shutting_down",
        "message": "Lead Intelligence server is shutting down. Port 8000 is being released.",
    }


# ── Static Files and Single-Page Dashboard ────────────────────────────────────

REACT_DIST_DIR = ROOT_DIR / "frontend-react" / "dist"
FRONTEND_DIR = REACT_DIST_DIR if REACT_DIST_DIR.exists() else (ROOT_DIR / "frontend")
ASSETS_DIR = FRONTEND_DIR / "assets"
CSS_DIR = FRONTEND_DIR / "css"
JS_DIR = FRONTEND_DIR / "js"

if ASSETS_DIR.exists():
    app.mount("/assets", StaticFiles(directory=str(ASSETS_DIR)), name="assets")
if CSS_DIR.exists():
    app.mount("/css", StaticFiles(directory=str(CSS_DIR)), name="css")
if JS_DIR.exists():
    app.mount("/js", StaticFiles(directory=str(JS_DIR)), name="js")
if FRONTEND_DIR.exists():
    app.mount("/static", StaticFiles(directory=str(FRONTEND_DIR)), name="static")


@app.get("/", response_class=HTMLResponse)
def serve_index():
    index_file = FRONTEND_DIR / "index.html"
    if not index_file.exists():
        return HTMLResponse("<h1>Lead Intelligence System</h1><p>Frontend assets not found.</p>", status_code=200)
    return HTMLResponse(
        index_file.read_text(encoding="utf-8"),
        headers={
            "Cache-Control": "no-cache, no-store, must-revalidate",
            "Pragma": "no-cache",
            "Expires": "0",
        },
    )


def run():
    """Convenience launcher for backend server."""
    import uvicorn
    uvicorn.run("backend.app.main:app", host="127.0.0.1", port=8000, reload=True)


if __name__ == "__main__":
    run()
