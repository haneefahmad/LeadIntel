"""
Business Lead Generator — XLSX exporter.
exporter.py: 2-sheet workbook with visual color legend and color-coded master records.

Sheets:
  1. Instructions      — visual column color legend and field reference guide
  2. Master_Database   — all lead records with color-coded headers, zebra striping, and frozen panes
"""

import logging
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

try:
    from backend.app import database as db
except ImportError:
    import database as db

logger = logging.getLogger(__name__)

# ── ARGB color constants ──────────────────────────────────────────────────────
YELLOW = "FFFBBF24"
INDIGO = "FF4F46E5"
ORANGE = "FFF59E0B"
GREEN  = "FF10B981"
TEAL   = "FF14B8A6"
PURPLE = "FF8B5CF6"
RED    = "FFEF4444"
GRAY   = "FF6B7280"
WHITE  = "FFFFFFFF"
DARK   = "FF111827"
BLUE   = "FF0B4CB8"
NAVY   = "FF16156C"


def _fill(argb: str) -> PatternFill:
    return PatternFill("solid", fgColor=argb)

def _font(bold: bool = False, color: str = WHITE, size: int = 10) -> Font:
    return Font(bold=bold, color=color, size=size)

def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _border() -> Border:
    s = Side(style="thin", color="FFD1D5DB")
    return Border(left=s, right=s, top=s, bottom=s)


# Column definitions follow database.MASTER_COLUMNS exactly.
_YELLOW = {"Record_ID", "Date_Added", "Added_By", "Lead_Status"}
_INDIGO = {"Company_Name", "Company_Type", "Primary_Industry", "Secondary_Industry", "Employee_Count"}
_BLUE   = {"Country", "State", "City", "Full_Address", "Postal_Code", "Google_Maps_URL"}
_ORANGE = {"Google_Rating", "Reviews_Count"}
_GREEN  = {"Primary_Phone", "General_Email", "WhatsApp_Number", "Website_URL", "Has_Website", "Company_LinkedIn"}
_TEAL   = {"DM_Full_Name", "DM_Title", "DM_Authority_Level", "DM_Direct_Email", "DM_Email_Status", "DM_Email_Score", "DM_Direct_Phone", "DM_LinkedIn_URL"}
_PURPLE = {"Pipeline_Stage", "Outreach_Status", "Assigned_To", "Deal_Value", "CRM_Notes", "Tags", "Last_Updated", "Updated_By"}


def _header_color(column: str) -> str:
    if column in _YELLOW:
        return YELLOW
    if column in _INDIGO:
        return INDIGO
    if column in _BLUE:
        return BLUE
    if column in _ORANGE:
        return ORANGE
    if column in _GREEN:
        return GREEN
    if column in _TEAL:
        return TEAL
    if column in _PURPLE:
        return PURPLE
    return INDIGO


MASTER_COLUMNS: list[tuple[str, str]] = [
    (column, _header_color(column)) for column in db.MASTER_COLUMNS
]


# ── Exporter ──────────────────────────────────────────────────────────────────

class XLSXExporter:

    def export(self, output_path: str, columns: list[str] | None = None) -> str:
        wb = Workbook()
        wb.remove(wb.active)

        self._instructions(wb)
        self._master_database(wb, columns=columns)

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        logger.info("XLSX exported → %s", output_path)
        return output_path

    # ── Sheet 1: Instructions (Color Legend & Field Reference) ───────────────
    def _instructions(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Instructions")
        ws.sheet_view.showGridLines = True

        # Set column widths
        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 44
        ws.column_dimensions["D"].width = 75

        # Title Block
        ws.merge_cells("A1:D1")
        title_cell = ws["A1"]
        title_cell.value = "Lead Intelligence Database — Column Color Legend"
        title_cell.font = Font(name="Calibri", bold=True, color=WHITE, size=14)
        title_cell.fill = _fill(NAVY)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 34

        # Subtitle Block
        ws.merge_cells("A2:D2")
        sub_cell = ws["A2"]
        sub_cell.value = "Visual reference guide explaining what each header color represents in the 'Master_Database' sheet."
        sub_cell.font = Font(name="Calibri", italic=True, color="FF475569", size=10)
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        ws.row_dimensions[3].height = 10  # Spacing row

        # Table Headers
        headers = ["Header Color", "Category / Data Type", "What It Represents", "Associated Columns in Master Database"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(4, col_idx, h)
            cell.font = Font(name="Calibri", bold=True, color=WHITE, size=10)
            cell.fill = _fill("FF1E293B")
            cell.alignment = _center()
            cell.border = _border()
        ws.row_dimensions[4].height = 26

        # Color rows definition
        legend_data = [
            (
                "YELLOW",
                YELLOW,
                "FF1F2937",
                "System & Identification",
                "Unique record IDs, ingestion date, campaign attribution, and overall lead qualification status.",
                "Record id, Date added, Added by, Lead status",
            ),
            (
                "INDIGO",
                INDIGO,
                WHITE,
                "Company Firmographics",
                "Standardized corporate entity name, legal structure, primary/secondary industries, and workforce size.",
                "Company Name, Company type, Primary Industry, Secondary Industry, Employee count",
            ),
            (
                "BLUE",
                BLUE,
                WHITE,
                "Location & Maps",
                "Municipal location, state jurisdiction, physical business address, postal zip code, and Google Maps URL.",
                "Country, State, City, Full Address, Postal code, Google maps url",
            ),
            (
                "ORANGE",
                ORANGE,
                WHITE,
                "Reputation & Reviews",
                "Public Google review total scores and customer review volume metrics.",
                "Google Rating, Reviews count",
            ),
            (
                "GREEN",
                GREEN,
                WHITE,
                "General Contact & Presence",
                "Corporate HQ switchboard, general contact inboxes, WhatsApp, official domain website, and LinkedIn company page.",
                "Primary phone, General Email, Whatsapp number, Website URL, Has Website, Company Linkedin",
            ),
            (
                "TEAL",
                TEAL,
                WHITE,
                "Decision Maker & Leadership",
                "Key executive contact details, verified corporate direct emails, verification scores, direct phone/mobile, and personal LinkedIn.",
                "DM Full name, DM Title, DM Authority Level, DM Direct Email, DM Email Status, DM Email Score, DM Direct phone, DM Linkedin URL",
            ),
            (
                "PURPLE",
                PURPLE,
                WHITE,
                "Pipeline & CRM",
                "Deal progress stages, outreach cadences, sales rep assignment, deal valuation, CRM notes, tags, and audit timestamps.",
                "Pipeline Stage, Outreach Status, Assigned to, Deal Value, CRM Notes, Tags, Last Updated, Updated by",
            ),
        ]

        for idx, (label, bg, fg, cat, desc, cols) in enumerate(legend_data, start=5):
            # Col A: Swatch
            c_color = ws.cell(idx, 1, label)
            c_color.font = Font(name="Calibri", bold=True, color=fg, size=10)
            c_color.fill = _fill(bg)
            c_color.alignment = _center()
            c_color.border = _border()

            # Col B: Category
            c_cat = ws.cell(idx, 2, cat)
            c_cat.font = Font(name="Calibri", bold=True, color="FF0F172A", size=10)
            c_cat.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            c_cat.border = _border()

            # Col C: Description
            c_desc = ws.cell(idx, 3, desc)
            c_desc.font = Font(name="Calibri", color="FF334155", size=9.5)
            c_desc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
            c_desc.border = _border()

            # Col D: Columns
            c_cols = ws.cell(idx, 4, cols)
            c_cols.font = Font(name="Consolas", color="FF1E293B", size=9)
            c_cols.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
            c_cols.border = _border()

            if idx % 2 == 0:
                c_cat.fill = _fill("FFF8FAFC")
                c_desc.fill = _fill("FFF8FAFC")
                c_cols.fill = _fill("FFF8FAFC")

            ws.row_dimensions[idx].height = 40

    # ── Sheet 2: Master_Database ──────────────────────────────────────────────
    def _master_database(self, wb: Workbook, columns: list[str] | None = None) -> None:
        ws = wb.create_sheet("Master_Database")
        ws.sheet_view.showGridLines = False

        if columns:
            active_cols = [(col, _header_color(col)) for col in columns if col in db.MASTER_COLUMNS]
            if not active_cols:
                active_cols = MASTER_COLUMNS
        else:
            active_cols = MASTER_COLUMNS

        freeze_col_idx = min(len(active_cols), 5)
        ws.freeze_panes = f"{get_column_letter(freeze_col_idx)}2"

        for col_idx, (col_name, color) in enumerate(active_cols, 1):
            display_title = db.COLUMN_DISPLAY_MAP.get(col_name, col_name)
            cell           = ws.cell(1, col_idx, display_title)
            cell.fill      = _fill(color)
            cell.font      = _font(bold=True)
            cell.alignment = _center()
            cell.border    = _border()
        ws.row_dimensions[1].height = 28

        records = db.get_all_records()
        for row_idx, record in enumerate(records, 2):
            for col_idx, (col_name, _) in enumerate(active_cols, 1):
                val  = record.get(col_name, "")
                cell = ws.cell(row_idx, col_idx, val if val is not None else "")
                cell.alignment = Alignment(vertical="center", wrap_text=False)
                cell.border    = _border()
                if row_idx % 2 == 0:
                    cell.fill = _fill("FFF9FAFB")
            ws.row_dimensions[row_idx].height = 18

        self._autofit(ws)

    # ── Utility ───────────────────────────────────────────────────────────────
    def _autofit(self, ws: Worksheet, min_w: int = 8, max_w: int = 40) -> None:
        for col_cells in ws.columns:
            col_letter = get_column_letter(col_cells[0].column)
            max_len    = max(
                (len(str(c.value)) for c in col_cells if c.value), default=0
            )
            ws.column_dimensions[col_letter].width = max(min_w, min(max_w, max_len + 2))
